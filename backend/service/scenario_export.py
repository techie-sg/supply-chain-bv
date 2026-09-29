"""Build an Excel workbook from a scenario without accessing the database."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import DateTime

from database.models import HourlyMetric, Order, Rider, Zone
from service.scenarios import TIMEZONE, build_scenario


def scenario_xlsx(key: str) -> bytes:
    rows, _ = build_scenario(key, datetime.now(TIMEZONE))
    workbook = Workbook()
    workbook.remove(workbook.active)

    for model in (Zone, Rider, Order, HourlyMetric):
        sheet = workbook.create_sheet(model.__tablename__)
        columns = list(model.__table__.columns)
        sheet.append(
            [
                f"{column.name}_ist"
                if isinstance(column.type, DateTime)
                else column.name
                for column in columns
            ]
        )
        for row in rows:
            if not isinstance(row, model):
                continue
            values = []
            for column in columns:
                value = getattr(row, column.key)
                if isinstance(value, datetime):
                    value = value.astimezone(TIMEZONE).replace(tzinfo=None)
                values.append(value)
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
