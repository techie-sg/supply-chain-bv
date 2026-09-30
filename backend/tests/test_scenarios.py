from datetime import datetime, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
import yaml
from openpyxl import load_workbook

import app as app_module
from blueprints import scenarios as scenario_routes
from database.models import HourlyMetric, Order, Rider, Zone
from service import scenarios
from service.scenarios import build_scenario, scenario_names


def test_every_scenario_builds_fresh_state() -> None:
    as_of = datetime(2026, 10, 3, 19, 30, tzinfo=ZoneInfo("Asia/Kolkata"))
    for key in (item["key"] for item in scenario_names()):
        rows, context = build_scenario(key, as_of)
        orders = [row for row in rows if isinstance(row, Order)]
        riders = [row for row in rows if isinstance(row, Rider)]
        history = [row for row in rows if isinstance(row, HourlyMetric)]
        zones = [row for row in rows if isinstance(row, Zone)]
        assert len(orders) == context["counts"]["orders"]
        assert len(riders) == 9 and len(history) == 12 and len(zones) == 3
        assert {row.date for row in history} == {
            as_of.date() - timedelta(days=1),
            as_of.date() - timedelta(days=2),
        }
        assert all(row.as_of == as_of for row in orders + riders)
        assert all((as_of - order.placed_at).total_seconds() > 0 for order in orders)
        assert {zone.zone_id for zone in zones} == {"Z-A", "Z-B", "Z-C"}
        assert "prompt" not in context


def test_normal_backlog_and_rain_have_distinct_starting_states() -> None:
    as_of = datetime(2026, 10, 3, 19, 30, tzinfo=ZoneInfo("Asia/Kolkata"))
    _, normal = build_scenario("normal", as_of)
    _, backlog = build_scenario("backlog", as_of)
    _, rain = build_scenario("rain", as_of)

    assert normal["counts"]["orders"] == 6
    assert normal["counts"]["available_riders"] == 4
    assert normal["is_raining"] is False
    assert backlog["counts"]["packed_waiting"] == 6
    assert backlog["counts"]["available_riders"] == 2
    assert backlog["is_raining"] is False
    assert rain["counts"]["packed_waiting"] == 8
    assert rain["is_raining"] is True
    assert len(rain["candidate_routes"]) == 8


def test_load_endpoint_uses_only_scenario_name(monkeypatch) -> None:
    called = []

    def fake_load(key):
        if key == "unknown":
            raise KeyError(key)
        called.append(key)
        return {"scenario_key": key}

    monkeypatch.setattr(scenario_routes, "load_scenario", fake_load)
    with app_module.app.test_client() as client:
        listing = client.get("/api/scenarios")
        assert listing.status_code == 200
        listing_data = listing.get_json()
        assert listing_data is not None
        assert {item["key"] for item in listing_data["scenarios"]} == {
            "normal",
            "backlog",
            "rain",
        }
        assert client.post("/api/scenarios/backlog/load").status_code == 200
        assert client.post("/api/scenarios/unknown/load").status_code == 404
        assert (
            client.post(
                "/api/scenarios/backlog/load",
                data='{"as_of":',
                content_type="application/json",
            ).status_code
            == 400
        )
        assert (
            client.post(
                "/api/scenarios/backlog/load",
                data="not JSON",
                content_type="text/plain",
            ).status_code
            == 400
        )
        assert (
            client.post(
                "/api/scenarios/backlog/load",
                data="null",
                content_type="application/json",
            ).status_code
            == 400
        )
        assert client.post("/api/scenarios/backlog/load", json={}).status_code == 400
        response = client.post("/api/scenarios/rain/load")
    assert response.status_code == 200
    assert called == ["backlog", "rain"]


def test_each_file_can_define_its_own_data(monkeypatch, tmp_path) -> None:
    source = scenarios.SCENARIO_DIR / "normal.yaml"
    data = yaml.safe_load(source.read_text())
    data["store_id"] = "TEST-STORE"
    data["zones"][0]["zone_id"] = "TEST-ZONE"
    for rider in data["riders"]:
        if rider["current_zone"] == "Z-A":
            rider["current_zone"] = "TEST-ZONE"
    for order in data["orders"]:
        if order["zone_id"] == "Z-A":
            order["zone_id"] = "TEST-ZONE"
    data["hourly_metrics"] = data["hourly_metrics"][:1]
    (tmp_path / "custom.yaml").write_text(yaml.safe_dump(data))
    monkeypatch.setattr(scenarios, "SCENARIO_DIR", tmp_path)

    rows, context = build_scenario(
        "custom", datetime(2026, 10, 3, 19, 30, tzinfo=ZoneInfo("Asia/Kolkata"))
    )
    assert scenario_names() == [{"key": "custom", "title": data["title"]}]
    assert context["store_id"] == "TEST-STORE"
    assert context["counts"]["hourly_metrics"] == 1
    assert any(isinstance(row, Zone) and row.zone_id == "TEST-ZONE" for row in rows)


def test_invalid_scenario_is_rejected_before_database_work(
    monkeypatch, tmp_path
) -> None:
    data = yaml.safe_load((scenarios.SCENARIO_DIR / "normal.yaml").read_text())
    data["orders"][0]["assigned_rider_id"] = "UNKNOWN"
    (tmp_path / "broken.yaml").write_text(yaml.safe_dump(data))
    monkeypatch.setattr(scenarios, "SCENARIO_DIR", tmp_path)

    with pytest.raises(ValueError, match="Invalid scenario file"):
        scenarios.load_scenario("broken")


def test_download_scenario_xlsx_without_loading_database(monkeypatch) -> None:
    def no_database(*args, **kwargs):
        raise AssertionError("download must not access the database")

    monkeypatch.setattr(scenarios, "get_session", no_database)
    with app_module.app.test_client() as client:
        response = client.get("/api/scenarios/normal/download")
        assert response.status_code == 200
        assert (
            response.headers["Content-Disposition"]
            == "attachment; filename=normal.xlsx"
        )
        assert response.mimetype == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert client.get("/api/scenarios/unknown/download").status_code == 404

    workbook = load_workbook(BytesIO(response.data), read_only=True)
    assert workbook.sheetnames == ["zones", "riders", "orders", "hourly_metrics"]
    assert [workbook[name].max_row - 1 for name in workbook.sheetnames] == [3, 9, 6, 12]
    orders = list(workbook["orders"].values)
    first = dict(zip(orders[0], orders[1], strict=True))
    assert first["order_id"] == "ORD-01-001"
    as_of = first["as_of_ist"]
    placed_at = first["placed_at_ist"]
    assert isinstance(as_of, datetime) and isinstance(placed_at, datetime)
    assert as_of - placed_at == timedelta(seconds=240)
